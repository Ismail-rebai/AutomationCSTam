"""
PowerPoint (.pptx) Presentation Compiler for OliveSoft Proposals.

Generates an executive-level, branded 16:9 widescreen presentation deck
incorporating technical solution, phase-by-step realization, timeline,
staffing roster, and transparent TND pricing breakdown.
"""

import os
from pathlib import Path
from typing import List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .proposal_generator import ProposalDeckData

# --- Brand Colors ---
COLOR_PRIMARY_DARK = RGBColor(38, 70, 43)    # Olive Dark (#26462B)
COLOR_PRIMARY = RGBColor(59, 110, 76)        # Forest Olive (#3B6E4C)
COLOR_ACCENT_GOLD = RGBColor(217, 119, 6)    # Amber Gold (#D97706)
COLOR_NAVY_DARK = RGBColor(15, 23, 42)       # Slate Navy (#0F172A)
COLOR_CARD_BG = RGBColor(248, 250, 252)      # Light Surface (#F8FAFC)
COLOR_CARD_BORDER = RGBColor(226, 232, 240)  # Border (#E2E8F0)
COLOR_TEXT_MAIN = RGBColor(30, 41, 59)       # Dark Slate (#1E293B)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)   # Muted Grey (#64748B)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_LIGHT_GREEN = RGBColor(236, 253, 245)  # Light Mint


def _add_header_banner(slide, title_text: str, category_text: str = "PROPOSITION TECHNIQUE & COMMERCIALE"):
    """Adds a standard OliveSoft header banner to content slides."""
    # Top banner background
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_PRIMARY_DARK
    top_bar.line.fill.background()

    # Category breadcrumb
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(10), Inches(0.3))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_ACCENT_GOLD

    # Main slide title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.40), Inches(11), Inches(0.6))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(20)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_WHITE

    # Logo / company tag in top right
    tag_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.35), Inches(2.2), Inches(0.5))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.alignment = PP_ALIGN.RIGHT
    p_tag.text = "OliveSoft AI"
    p_tag.font.size = Pt(14)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_ACCENT_GOLD


def _add_footer(slide, current_slide: int, total_slides: int = 8):
    """Adds confidential footer and slide number."""
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(8), Inches(0.35))
    p = footer_box.text_frame.paragraphs[0]
    p.text = "OliveSoft Engineering — Confidentiel — Document d'Offre Commerciale & Technique"
    p.font.size = Pt(8.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    num_box = slide.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.2), Inches(0.35))
    p2 = num_box.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    p2.text = f"{current_slide} / {total_slides}"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_MUTED


# --- Slide 1: Title & Executive Summary ---

def _build_slide_1_title(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Background header block (dark olive)
    header_block = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(4.3))
    header_block.fill.solid()
    header_block.fill.fore_color.rgb = COLOR_PRIMARY_DARK
    header_block.line.fill.background()

    # Company & Subject Tag
    tag_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.6), Inches(11.5), Inches(0.4))
    p_tag = tag_box.text_frame.paragraphs[0]
    p_tag.text = f"APPEL D'OFFRES N° {data.tender_id} — OFFRE TECHNIQUE & FINANCIÈRE"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_ACCENT_GOLD

    # Main Proposal Title
    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.1), Inches(11.5), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p_title = tf.paragraphs[0]
    p_title.text = data.tender_title
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    # Client subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.7), Inches(11.5), Inches(0.6))
    p_sub = sub_box.text_frame.paragraphs[0]
    p_sub.text = f"Préparé par OliveSoft pour : {data.client_name}"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = RGBColor(209, 250, 229)

    # 3 Executive Metric Highlight Cards
    card_width = Inches(3.6)
    card_height = Inches(2.1)
    card_y = Inches(4.65)

    metrics = [
        ("DURÉE GLOBALE DU PROJET", f"{data.total_duration_weeks} Semaines", "5 Phases d'exécution ordonnancées"),
        ("BUDGET TOTAL ESTIMÉ", f"{data.pricing.total_price_tnd:,.2f} TND HT", f"Inclut buffer de contingence ({data.pricing.contingency_percent:.0f}%)"),
        ("INDICE DE COUVERTURE RAG", f"{int(round(data.fit_score * 100))}% Fit Score", "Adéquation prouvée aux compétences OliveSoft"),
    ]

    for i, (m_title, m_val, m_sub) in enumerate(metrics):
        card_x = Inches(0.9 + i * 3.9)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_PRIMARY
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(card_x + Inches(0.2), card_y + Inches(0.2), card_width - Inches(0.4), card_height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = m_title
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_MUTED

        p2 = tf.add_paragraph()
        p2.text = m_val
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_PRIMARY_DARK

        p3 = tf.add_paragraph()
        p3.text = m_sub
        p3.font.size = Pt(9)
        p3.font.color.rgb = COLOR_TEXT_MAIN

    _add_footer(slide, 1)


# --- Slide 2: Context, Challenges & Business Objectives ---

def _build_slide_2_context(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header_banner(slide, "Compréhension du Contexte & Objectifs Stratégiques", "ANALYSE DU BESOIN")

    col_width = Inches(5.6)
    col_height = Inches(5.3)

    # Left Column: Pain points
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.45), col_width, col_height)
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = COLOR_CARD_BG
    left_card.line.color.rgb = COLOR_CARD_BORDER

    tb_left = slide.shapes.add_textbox(Inches(1.1), Inches(1.65), col_width - Inches(0.4), col_height - Inches(0.4))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "Défis & Problématiques Identifiés"
    p_lh.font.size = Pt(16)
    p_lh.font.bold = True
    p_lh.font.color.rgb = COLOR_NAVY_DARK

    for point in data.client_pain_points:
        p = tf_l.add_paragraph()
        p.text = f"•  {point}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_before = Pt(12)

    # Right Column: Business Objectives
    right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.45), col_width, col_height)
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = COLOR_LIGHT_GREEN
    right_card.line.color.rgb = COLOR_PRIMARY

    tb_right = slide.shapes.add_textbox(Inches(7.0), Inches(1.65), col_width - Inches(0.4), col_height - Inches(0.4))
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "Objectifs Métier & Bénéfices Attendus"
    p_rh.font.size = Pt(16)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_PRIMARY_DARK

    for obj in data.business_objectives:
        p = tf_r.add_paragraph()
        p.text = f"✔  {obj}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_before = Pt(12)

    _add_footer(slide, 2)


# --- Slide 3: Technical Architecture & Stack ---

def _build_slide_3_architecture(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header_banner(slide, "Architecture Technique Cible & Stack Recommandée", "SOLUTION TECHNIQUE")

    # Paradigm Banner
    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.4), Inches(11.5), Inches(1.1))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_PRIMARY_DARK
    banner.line.fill.background()

    tb_b = slide.shapes.add_textbox(Inches(1.1), Inches(1.45), Inches(11.1), Inches(1.0))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p_b1 = tf_b.paragraphs[0]
    p_b1.text = data.technical_solution.architecture_style.upper()
    p_b1.font.size = Pt(14)
    p_b1.font.bold = True
    p_b1.font.color.rgb = COLOR_ACCENT_GOLD

    p_b2 = tf_b.add_paragraph()
    p_b2.text = data.technical_solution.summary
    p_b2.font.size = Pt(10.5)
    p_b2.font.color.rgb = COLOR_WHITE

    # 4 Tech Stack Cards
    stack_items = list(data.technical_solution.tech_stack.items())
    card_w = Inches(2.7)
    card_h = Inches(4.0)

    for i, (category, tools) in enumerate(stack_items[:4]):
        card_x = Inches(0.9 + i * 2.93)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, Inches(2.75), card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER

        tb = slide.shapes.add_textbox(card_x + Inches(0.15), Inches(2.85), card_w - Inches(0.3), card_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = category
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_PRIMARY_DARK

        for tool in tools:
            p = tf.add_paragraph()
            p.text = f"• {tool}"
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_before = Pt(6)

    _add_footer(slide, 3)


# --- Slide 4: Step-by-Step Realization (5 Phases) ---

def _build_slide_4_phases(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header_banner(slide, "Plan de Réalisation Pas-à-Pas (5 Phases)", "MÉTHODOLOGIE D'EXÉCUTION")

    card_w = Inches(2.2)
    card_h = Inches(5.3)

    for i, phase in enumerate(data.phases[:5]):
        card_x = Inches(0.85 + i * 2.36)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, Inches(1.45), card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_PRIMARY if i % 2 == 0 else COLOR_CARD_BORDER

        # Phase header ribbon
        ribbon = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, Inches(1.45), card_w, Inches(0.85))
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = COLOR_PRIMARY_DARK if i % 2 == 0 else COLOR_PRIMARY
        ribbon.line.fill.background()

        tb_rib = slide.shapes.add_textbox(card_x + Inches(0.1), Inches(1.5), card_w - Inches(0.2), Inches(0.75))
        tf_r = tb_rib.text_frame
        tf_r.word_wrap = True
        p_num = tf_r.paragraphs[0]
        p_num.text = f"PHASE {phase.phase_num} ({phase.duration_weeks} SEM.)"
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT_GOLD

        p_name = tf_r.add_paragraph()
        p_name.text = phase.name
        p_name.font.size = Pt(10)
        p_name.font.bold = True
        p_name.font.color.rgb = COLOR_WHITE

        # Content
        tb_body = slide.shapes.add_textbox(card_x + Inches(0.12), Inches(2.4), card_w - Inches(0.24), card_h - Inches(1.0))
        tf_b = tb_body.text_frame
        tf_b.word_wrap = True

        p_del = tf_b.paragraphs[0]
        p_del.text = "Livrables clés :"
        p_del.font.size = Pt(9.5)
        p_del.font.bold = True
        p_del.font.color.rgb = COLOR_NAVY_DARK

        for d in phase.deliverables[:3]:
            p = tf_b.add_paragraph()
            p.text = f"• {d}"
            p.font.size = Pt(8.5)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_before = Pt(4)

        p_act = tf_b.add_paragraph()
        p_act.text = "Activités :"
        p_act.font.size = Pt(9.5)
        p_act.font.bold = True
        p_act.font.color.rgb = COLOR_NAVY_DARK
        p_act.space_before = Pt(8)

        for a in phase.key_activities[:2]:
            p = tf_b.add_paragraph()
            p.text = f"- {a}"
            p.font.size = Pt(8.5)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_before = Pt(3)

    _add_footer(slide, 4)


# --- Slide 5: Roadmap & Governance ---

def _build_slide_5_roadmap(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header_banner(slide, "Planning des Jalons & Gouvernance du Projet", "PLANNING & PILOTAGE")

    # Table of Milestones
    table_shape = slide.shapes.add_table(6, 4, Inches(0.9), Inches(1.5), Inches(11.5), Inches(3.2))
    table = table_shape.table

    # Column widths
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(2.5)
    table.columns[3].width = Inches(3.0)

    headers = ["Jalon", "Phase & Activité Clé", "Échéance", "Livrable Validé"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    curr_week = 0
    for row_idx, phase in enumerate(data.phases[:5], start=1):
        start_w = curr_week + 1
        curr_week += phase.duration_weeks
        end_w = curr_week

        row_data = [
            f"M{row_idx}",
            f"{phase.name}",
            f"Semaines {start_w} à {end_w}",
            phase.deliverables[0] if phase.deliverables else "Validation Phase",
        ]
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_BG if row_idx % 2 == 0 else COLOR_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_TEXT_MAIN

    # Governance Card below
    gov_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(4.9), Inches(11.5), Inches(1.9))
    gov_card.fill.solid()
    gov_card.fill.fore_color.rgb = COLOR_LIGHT_GREEN
    gov_card.line.color.rgb = COLOR_PRIMARY

    tb_gov = slide.shapes.add_textbox(Inches(1.1), Inches(5.0), Inches(11.1), Inches(1.7))
    tf_g = tb_gov.text_frame
    tf_g.word_wrap = True

    p_gh = tf_g.paragraphs[0]
    p_gh.text = "Modèle de Gouvernance & Rituels Agiles OliveSoft"
    p_gh.font.size = Pt(13)
    p_gh.font.bold = True
    p_gh.font.color.rgb = COLOR_PRIMARY_DARK

    points = [
        "Comité de Pilotage (COPIL) mensuel : Arbitrages stratégiques, validation des jalons et suivi budgétaire.",
        "Comité Technique (COTEC) bi-hebdomadaire : Suivi opérationnel des sprints, démos de fonctionnalités et levée des bloquants.",
        "Sprints de 2 semaines avec Daily Standup et rétrospectives d'amélioration continue.",
    ]
    for pt in points:
        p = tf_g.add_paragraph()
        p.text = f"✔ {pt}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_TEXT_MAIN
        p.space_before = Pt(4)

    _add_footer(slide, 5)


# --- Slide 6: Dedicated Project Team ---

def _build_slide_6_team(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header_banner(slide, "Équipe Projet Dédiée (Roster Mobilisé)", "RESSOURCES HUMAINES")

    card_w = Inches(3.6)
    card_h = Inches(5.3)

    for i, profile in enumerate(data.staffing[:3]):
        card_x = Inches(0.9 + i * 3.9)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, Inches(1.45), card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_PRIMARY

        tb = slide.shapes.add_textbox(card_x + Inches(0.2), Inches(1.65), card_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_name = tf.paragraphs[0]
        p_name.text = profile.name
        p_name.font.size = Pt(17)
        p_name.font.bold = True
        p_name.font.color.rgb = COLOR_PRIMARY_DARK

        p_role = tf.add_paragraph()
        p_role.text = profile.role_in_project
        p_role.font.size = Pt(11)
        p_role.font.bold = True
        p_role.font.color.rgb = COLOR_ACCENT_GOLD
        p_role.space_before = Pt(4)

        p_t = tf.add_paragraph()
        p_t.text = f"Titre : {profile.title}"
        p_t.font.size = Pt(10)
        p_t.font.color.rgb = COLOR_TEXT_MUTED
        p_t.space_before = Pt(8)

        p_cv = tf.add_paragraph()
        p_cv.text = f"Référence CV : {profile.cv_id} (Certifié OliveSoft)"
        p_cv.font.size = Pt(9.5)
        p_cv.font.color.rgb = COLOR_TEXT_MUTED

        p_sk_h = tf.add_paragraph()
        p_sk_h.text = "Compétences clés validées :"
        p_sk_h.font.size = Pt(11)
        p_sk_h.font.bold = True
        p_sk_h.font.color.rgb = COLOR_NAVY_DARK
        p_sk_h.space_before = Pt(14)

        for sk in profile.matched_skills[:5]:
            p = tf.add_paragraph()
            p.text = f"• {sk}"
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_TEXT_MAIN
            p.space_before = Pt(3)

    _add_footer(slide, 6)


# --- Slide 7: Commercial Proposal & Price Breakdown (TND) ---

def _build_slide_7_pricing(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header_banner(slide, "Offre Financière & Décomposition Budgétaire (TND)", "PROPOSITION COMMERCIALE")

    # Table of Roles
    roles = data.pricing.roles
    total_rows = len(roles) + 4  # Header + Roles + Subtotal + Contingency + TOTAL
    table_shape = slide.shapes.add_table(total_rows, 4, Inches(0.9), Inches(1.45), Inches(8.2), Inches(5.3))
    table = table_shape.table

    table.columns[0].width = Inches(3.6)
    table.columns[1].width = Inches(1.5)
    table.columns[2].width = Inches(1.3)
    table.columns[3].width = Inches(1.8)

    headers = ["Profil & Rôle d'Expertise", "Taux J/H (TND)", "Jours-H (J/H)", "Total HT (TND)"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    # Role rows
    for row_idx, r in enumerate(roles, start=1):
        row_data = [
            r.role,
            f"{r.rate_per_day_tnd:,.0f} TND",
            f"{r.days} j/h",
            f"{r.total_tnd:,.2f} TND",
        ]
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_BG if row_idx % 2 == 0 else COLOR_WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(9.5)
            p.font.color.rgb = COLOR_TEXT_MAIN

    # Subtotal
    idx_sub = len(roles) + 1
    table.cell(idx_sub, 0).text = "Sous-total Prestations d'Ingénierie"
    table.cell(idx_sub, 0).text_frame.paragraphs[0].font.bold = True
    table.cell(idx_sub, 3).text = f"{data.pricing.base_cost_tnd:,.2f} TND"
    table.cell(idx_sub, 3).text_frame.paragraphs[0].font.bold = True

    # Contingency
    idx_cont = len(roles) + 2
    table.cell(idx_cont, 0).text = f"Provision pour Risques & Imprévus ({data.pricing.contingency_percent:.0f}%)"
    table.cell(idx_cont, 3).text = f"{data.pricing.contingency_tnd:,.2f} TND"

    # TOTAL ROW (highlighted)
    idx_tot = len(roles) + 3
    for c in range(4):
        table.cell(idx_tot, c).fill.solid()
        table.cell(idx_tot, c).fill.fore_color.rgb = COLOR_PRIMARY

    table.cell(idx_tot, 0).text = "TOTAL FORFAIT PROPOSITION (TND HT)"
    p_tot0 = table.cell(idx_tot, 0).text_frame.paragraphs[0]
    p_tot0.font.bold = True
    p_tot0.font.size = Pt(11)
    p_tot0.font.color.rgb = COLOR_WHITE

    table.cell(idx_tot, 3).text = f"{data.pricing.total_price_tnd:,.2f} TND"
    p_tot3 = table.cell(idx_tot, 3).text_frame.paragraphs[0]
    p_tot3.font.bold = True
    p_tot3.font.size = Pt(12)
    p_tot3.font.color.rgb = COLOR_ACCENT_GOLD

    # Right Card: SLA Maintenance & Terms
    card_sla = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.3), Inches(1.45), Inches(3.1), Inches(5.3))
    card_sla.fill.solid()
    card_sla.fill.fore_color.rgb = COLOR_LIGHT_GREEN
    card_sla.line.color.rgb = COLOR_PRIMARY

    tb_sla = slide.shapes.add_textbox(Inches(9.45), Inches(1.65), Inches(2.8), Inches(4.9))
    tf_s = tb_sla.text_frame
    tf_s.word_wrap = True

    p_sh = tf_s.paragraphs[0]
    p_sh.text = "Garantie & Support SLA"
    p_sh.font.size = Pt(13)
    p_sh.font.bold = True
    p_sh.font.color.rgb = COLOR_PRIMARY_DARK

    p_g = tf_s.add_paragraph()
    p_g.text = "• Garantie 90 jours offerte : Prise en charge intégrale des anomalies post-lancement."
    p_g.font.size = Pt(9.5)
    p_g.space_before = Pt(8)

    p_sla_val = tf_s.add_paragraph()
    p_sla_val.text = f"• Maintenance SLA annuelle (Option) :\n  {data.pricing.maintenance_sla_annual_tnd:,.2f} TND HT / an (Support N2/N3, correctifs de sécurité, astreinte 4h)."
    p_sla_val.font.size = Pt(9.5)
    p_sla_val.space_before = Pt(10)

    p_terms = tf_s.add_paragraph()
    p_terms.text = "Conditions de paiement :\n- 30% au démarrage (Phase 1)\n- 30% recette intermédiaire (Phase 3)\n- 30% recette UAT (Phase 4)\n- 10% procès-verbal définitif."
    p_terms.font.size = Pt(8.5)
    p_terms.font.color.rgb = COLOR_TEXT_MUTED
    p_terms.space_before = Pt(12)

    _add_footer(slide, 7)


# --- Slide 8: Risks, Commitments & Next Steps ---

def _build_slide_8_risks(prs: Presentation, data: ProposalDeckData):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header_banner(slide, "Matrice des Risques, Engagements & Prochaines Étapes", "ENGAGEMENTS & CLÔTURE")

    # Left: Risks table / cards
    left_w = Inches(7.0)
    left_h = Inches(5.3)

    card_risk = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.45), left_w, left_h)
    card_risk.fill.solid()
    card_risk.fill.fore_color.rgb = COLOR_CARD_BG
    card_risk.line.color.rgb = COLOR_CARD_BORDER

    tb_r = slide.shapes.add_textbox(Inches(1.1), Inches(1.6), left_w - Inches(0.4), left_h - Inches(0.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "Maîtrise des Risques Opérationnels & Techniques"
    p_rh.font.size = Pt(14)
    p_rh.font.bold = True
    p_rh.font.color.rgb = COLOR_NAVY_DARK

    for r_item in data.risk_mitigations[:3]:
        p1 = tf_r.add_paragraph()
        p1.text = f"Risque : {r_item.risk}"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK
        p1.space_before = Pt(10)

        p2 = tf_r.add_paragraph()
        p2.text = f"Mitigation : {r_item.mitigation}"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_MAIN
        p2.space_before = Pt(2)

    # Right: Commitments & Contact
    right_w = Inches(4.3)
    card_comm = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(1.45), right_w, left_h)
    card_comm.fill.solid()
    card_comm.fill.fore_color.rgb = COLOR_PRIMARY_DARK
    card_comm.line.fill.background()

    tb_c = slide.shapes.add_textbox(Inches(8.3), Inches(1.6), right_w - Inches(0.4), left_h - Inches(0.3))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p_ch = tf_c.paragraphs[0]
    p_ch.text = "Engagements OliveSoft"
    p_ch.font.size = Pt(14)
    p_ch.font.bold = True
    p_ch.font.color.rgb = COLOR_ACCENT_GOLD

    comms = [
        "Cession totale des droits de propriété intellectuelle (code source, DAT, SFD).",
        "Disponibilité immédiate de l'équipe dès signature du contrat.",
        "Audit de sécurité indépendant avant mise en production.",
        "Assurance responsabilité civile professionnelle conforme aux exigences du cahier des charges.",
    ]
    for c in comms:
        p = tf_c.add_paragraph()
        p.text = f"✔ {c}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(8)

    p_contact = tf_c.add_paragraph()
    p_contact.text = "Contacts Équipe Commerciale :\nEmail : contact@olivesoft.tn\nTél : +216 71 000 000\nTunis, Tunisie"
    p_contact.font.size = Pt(9.5)
    p_contact.font.bold = True
    p_contact.font.color.rgb = RGBColor(209, 250, 229)
    p_contact.space_before = Pt(16)

    _add_footer(slide, 8)


# --- Master Presentation Builder ---

def build_proposal_presentation(data: ProposalDeckData, output_path: str) -> str:
    """
    Compiles a complete 8-slide PowerPoint presentation (.pptx)
    from structured ProposalDeckData.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _build_slide_1_title(prs, data)
    _build_slide_2_context(prs, data)
    _build_slide_3_architecture(prs, data)
    _build_slide_4_phases(prs, data)
    _build_slide_5_roadmap(prs, data)
    _build_slide_6_team(prs, data)
    _build_slide_7_pricing(prs, data)
    _build_slide_8_risks(prs, data)

    out_file = Path(output_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_file))

    return str(out_file)
